#!/usr/bin/env python3
"""Generate the launch project plan as both LAUNCH_PLAN.xlsx and LAUNCH_PLAN.md.

Single source of truth = the TASKS list below. Edit a row, re-run, both files
regenerate. Columns: ID, Epic, Task, Detail, Persona, Priority, Size, Deps, Status.

Personas: Dev=Developer/PM, Super=Sims super user, End=Sims end user.
Priority: P0=go-live gate, P1=want at launch, P2=fast-follow, GW=v2 groundwork (do at launch), V2=marketplace v2.
Size: S/M/L.
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

# (ID, Epic, Task, Detail, Persona, Priority, Size, Deps)
TASKS = [
    # --- Epic A: Account recovery + email (the long pole) ---
    ("A1", "Account recovery + email", "Choose email provider", "Pick Resend / Postmark / SES. Decide on sending domain + DKIM/SPF setup.", "Dev", "P0", "S", ""),
    ("A2", "Account recovery + email", "password_reset_tokens table", "Idempotent migration: id, user_id, token_hash, expires_at, used_at, created_at. Index (user_id),(token_hash).", "Dev", "P0", "S", ""),
    ("A3", "Account recovery + email", "POST /auth/forgot", "Issue hashed single-use token w/ expiry; always 200 (no account enumeration); send reset email.", "Dev", "P0", "M", "A1,A2"),
    ("A4", "Account recovery + email", "POST /auth/reset", "Validate token (unexpired, unused), set new password hash, invalidate token + existing sessions.", "Dev", "P0", "M", "A2"),
    ("A5", "Account recovery + email", "Forgot + reset UI pages", "Public 'forgot password' form and token-landing 'set new password' page; success/expired states.", "End", "P0", "M", "A3,A4"),
    ("A6", "Account recovery + email", "Transactional email templates", "Branded reset email (+ verify email template). Plain-text fallback. Link expiry copy.", "Dev", "P0", "S", "A1"),
    ("A7", "Account recovery + email", "Rate-limit forgot/reset", "Per-IP + per-account throttle on forgot/reset to prevent abuse + email bombing.", "Dev", "P0", "S", "A3,A4"),
    ("A8", "Account recovery + email", "Email verification on signup", "Store email_verified; send verify email; nudge (don't hard-block) unverified users.", "Dev", "P1", "M", "A1,A6"),

    # --- Epic B: Anonymous landing + featured (also 'login clarity') ---
    ("B1", "Anonymous landing", "Public root route /", "Make '/' a public landing (no auth redirect). Keep planner routes gated behind the auth wall.", "Dev", "P0", "S", ""),
    ("B2", "Anonymous landing", "Landing hero + pitch", "Hero copy, product screenshots, clear signup/login CTAs. Conveys plan + track + show off.", "End", "P0", "M", "B1"),
    ("B3", "Anonymous landing", "Featured saves strip", "Pull N hand-picked public showcases (/share/:token) as cards on the landing. Grows into the v2 directory.", "End", "P0", "M", "B1"),
    ("B4", "Anonymous landing", "Featured-picker config", "Config/admin way to set which showcase tokens are featured.", "Dev", "P0", "S", "B3"),
    ("B5", "Anonymous landing", "Login/signup clarity", "Distinct sign-up vs log-in CTAs + copy; clear mode toggle. (Your P0 'login clarity' item.)", "End", "P0", "S", "B1"),
    ("B6", "Anonymous landing", "Public nav + footer", "Public-zone header/footer linking About / Help / Privacy / ToS, reachable without an account.", "Dev", "P0", "S", "B1"),

    # --- Epic C: Legal layer ---
    ("C1", "Legal", "Privacy Policy page", "What data is stored (account, saves, photos, URLs), retention, deletion, third parties (Sentry, email).", "Dev", "P0", "M", ""),
    ("C2", "Legal", "Terms of Service page", "Acceptable use, UGC ownership, takedown, no-warranty, governing law.", "Dev", "P0", "M", ""),
    ("C3", "Legal", "Expand EA disclaimer + donation note", "Strengthen fan-content disclaimer; explicitly state donation-only, non-commercial (helps EA-policy posture).", "Dev", "P0", "S", ""),
    ("C4", "Legal", "Cookie/analytics notice", "Only if analytics/cookies added. Lightweight notice + link to privacy.", "Dev", "P0", "S", ""),
    ("C5", "Legal", "Wire legal links + signup consent", "Footer links + signup checkbox/consent referencing Privacy + ToS.", "Dev", "P0", "S", "C1,C2"),

    # --- Epic D: SSRF hardening ---
    ("D1", "Security (SSRF)", "Audit user-URL fetch points", "Enumerate every place server fetches a user-supplied URL (saveFiles download link, preview image, etc.).", "Dev", "P0", "S", ""),
    ("D2", "Security (SSRF)", "SSRF guard util", "Resolve + reject private/internal/metadata ranges (127/8,10/8,172.16/12,192.168/16,169.254/16,::1,localhost).", "Dev", "P0", "M", "D1"),
    ("D3", "Security (SSRF)", "Scheme allowlist + limits", "https-only; cap redirects; cap response size + timeout to prevent abuse.", "Dev", "P0", "S", "D2"),
    ("D4", "Security (SSRF)", "Apply guard everywhere", "Route all outbound user-URL fetches through the guard.", "Dev", "P0", "S", "D2,D3"),
    ("D5", "Security (SSRF)", "SSRF guard tests", "Assert guard rejects metadata/loopback/private IPs and DNS-rebind-style hosts.", "Dev", "P0", "S", "D4"),

    # --- Epic E: Data-loss safety net ---
    ("E1", "Data safety", "Verify backup round-trips", "Export a save via SaveSettings 'Download backup' -> re-import -> assert identical (lots/hh/sims/photos).", "Dev", "P0", "M", ""),
    ("E2", "Data safety", "Soft-delete saves", "deleted_at column; delete = soft; queries exclude soft-deleted; purge job later.", "Dev", "P0", "M", ""),
    ("E3", "Data safety", "Restore / undo window", "UI to restore a recently deleted save within a grace window.", "Super", "P1", "M", "E2"),
    ("E4", "Data safety", "Snapshot before destructive import", "Auto-snapshot/export before an overwrite import so a bad import is recoverable.", "Dev", "P1", "S", "E1"),

    # --- Epic F: Mobile triage ---
    ("F1", "Mobile", "Define must-work-on-mobile set", "Lock the in-scope mobile views: dashboard, households, lots, photos, public showcase. Edit = desktop-first.", "Dev", "P0", "S", ""),
    ("F2", "Mobile", "Responsive audit", "Audit in-scope views at 375 + 768 widths; catalogue breakages.", "End", "P0", "M", "F1"),
    ("F3", "Mobile", "Mobile nav (drawer)", "Sidebar -> off-canvas drawer / bottom nav on small screens.", "End", "P0", "M", "F2"),
    ("F4", "Mobile", "Fix in-scope layouts", "Make dashboard, lists, photos, showcase usable on mobile.", "End", "P0", "L", "F2"),
    ("F5", "Mobile", "Graceful desktop-only editors", "On mobile, present a clear 'best on desktop' state for out-of-scope editors instead of broken UI.", "End", "P1", "S", "F1"),

    # --- Epic G: Observability & ops ---
    ("G1", "Observability", "Confirm prod Sentry DSNs", "Verify VITE_SENTRY_DSN (browser) + server DSN set in prod env. Code already wired.", "Dev", "P1", "S", ""),
    ("G2", "Observability", "Source map upload", "Upload source maps in build so prod stack traces are readable.", "Dev", "P1", "S", "G1"),
    ("G3", "Observability", "Error-spike alert", "One alert rule routing error spikes to a channel/email.", "Dev", "P1", "S", "G1"),
    ("G4", "Observability", "Auth rate-limiting (shared)", "Generalize A7 throttle across login/signup/public endpoints.", "Dev", "P1", "S", "A7"),
    ("G5", "Observability", "Funnel analytics", "Track signup -> first save -> first household to measure onboarding.", "Dev", "P2", "M", ""),

    # --- Epic H: Parser / data correctness ---
    ("H1", "Parser correctness", "Expand parser regression tests", "Lock clubs/businesses/sims/lot-types against fixture saves so format work can't silently regress.", "Dev", "P1", "M", ""),
    ("H2", "Parser correctness", "PT-locale club names", "Fix Portuguese club name/description drop, OR document as a known limitation in Help.", "Dev", "P1", "M", ""),
    ("H3", "Parser correctness", "Large-save perf pass", "Profile import + render on a fully-built save (many lots/sims); fix hotspots.", "Dev", "P2", "M", ""),

    # --- Epic I: Super-user parity ---
    ("I1", "Super-user parity", "Notes vs descriptions schema", "Add private note + public description fields per the locked section matrix.", "Dev", "P1", "M", ""),
    ("I2", "Super-user parity", "World showcase blurb field", "Manual public blurb for worlds (no in-game field exists).", "Super", "P1", "S", "I1"),
    ("I3", "Super-user parity", "Notes/descriptions in editors", "Surface note + description inputs per section per the matrix.", "Super", "P1", "M", "I1"),
    ("I4", "Super-user parity", "Descriptions public / notes private", "Show descriptions on showcase; never expose notes.", "Super", "P1", "S", "I1"),
    ("I5", "Super-user parity", "Traits/aspirations on Sims roster", "Add trait + aspiration filters/columns to the Sims roster (data already parsed).", "Super", "P1", "M", ""),
    ("I6", "Super-user parity", "Migration IN (CSV / bulk entry)", "Let spreadsheet users import or fast-enter households/sims without re-typing. Biggest 'will I switch' lever.", "Super", "P1", "L", ""),
    ("I7", "Super-user parity", "Full data export (JSON)", "Account/save export so power users can always get their data out.", "Super", "P2", "M", ""),

    # --- Epic J: End-user delight & onboarding ---
    ("J1", "Delight & onboarding", "Magic-moment onboarding", "Lead new users with '.save import -> watch your world appear'. The wow demo.", "End", "P1", "M", ""),
    ("J2", "Delight & onboarding", "From-scratch quickstart", "Starter templates / pre-seeded empty Newcrest for the no-save crowd.", "End", "P1", "M", ""),
    ("J3", "Delight & onboarding", "Surface randomizer + inspo early", "Put the fun features front-of-journey.", "End", "P1", "S", ""),
    ("J4", "Delight & onboarding", "Discoverable share links", "Make 'share my save' obvious from inside the planner.", "End", "P1", "S", ""),
    ("J5", "Delight & onboarding", "Per-world progress rings", "'12 of 27 lots built' completion ring on the dashboard.", "End", "P2", "S", ""),
    ("J6", "Delight & onboarding", "One-tap 'surprise me' roll", "Combined world -> lot -> type roll for instant inspiration.", "End", "P2", "S", ""),
    ("J7", "Delight & onboarding", "Before/after lot cards", "Inspo vs built photo comparison on lot cards.", "End", "P2", "M", ""),
    ("J8", "Delight & onboarding", "Shareable image cards", "Export world/household cards as images for Reddit/Discord.", "End", "P2", "M", ""),
    ("J9", "Delight & onboarding", "Redo favicons (current set is ugly)", "Replace the favicon set (favicon.ico/.svg/-16/-32 + apple-touch-icon) with a polished, on-brand mark. Assets are already wired in index.html + the guide <head> (scripts/seo-head.mjs) — just swap the files in public/. Consider matching the plumbob/OG brand mark.", "End", "P1", "S", ""),

    # --- Epic CP: Creator preview / GTM (V1 must-do; the marketplace SUPPLY wedge) ---
    # Show ~10-20 hand-picked mega-creators the TOOL on their own save (read-only),
    # not just the public showcase. Earns trust + a feedback loop BEFORE the V2
    # marketplace. Curated/manual = no public-upload theft surface yet. See
    # project_creator_preview memory.
    ("CP1", "Creator preview", "Read-only preview data access", "A preview/share token that exposes a save's FULL planner dataset read-only, no auth: lots, households, sims (traits/occult/lifestage), clubs, businesses, diversity inputs. Extends the existing public showcase endpoints.", "Dev", "P1", "M", ""),
    ("CP2", "Creator preview", "Read-only planner preview UI", "Render the real planner backend read-only from a preview token, no signup: dashboard + world view + Diversity (the key wow). 'Preview' chrome; navigation works, edits disabled.", "End", "P1", "L", "CP1"),
    ("CP3", "Creator preview", "Conversion CTA / funnel", "From the preview: 'Sign up + drop your .save to make it yours (editable)' — convert creators (who already have their .save) in ~30s.", "End", "P1", "S", "CP2"),
    ("CP4", "Creator preview", "Per-creator import + preview link", "Admin/manual flow to import a creator's published .save and mint a preview link to DM. Low volume (10-20), manual is fine.", "Dev", "P1", "M", "CP1"),
    ("CP5", "Creator preview", "Read-only editor preview", "Show write-tools (lot editor, inspo) in a non-editable demo state — where 'the ease of the work' really lands. Needs a design pass (no-op buttons vs 'sign up to edit' nudge vs watermark).", "End", "P2", "M", "CP2"),

    # --- Epic K: Marketplace (groundwork now, build v2) ---
    ("K1", "Marketplace", "'List publicly' opt-in flag", "GROUNDWORK AT LAUNCH: add the publish flag + showcase data model so v2 is a query, not a rewrite.", "Dev", "GW", "S", ""),
    ("K2", "Marketplace", "Listing metadata model", "GROUNDWORK: cover image, last-updated, pack-requirements derived from parsed pack usage.", "Dev", "GW", "M", "K1"),
    ("K3", "Marketplace", "Directory page", "Anonymous marketplace directory aggregating public listings.", "End", "V2", "L", "K2"),
    ("K4", "Marketplace", "Search / sort / filter", "latest, popular, by world, by pack requirements, occult, lot count.", "End", "V2", "L", "K3"),
    ("K5", "Marketplace", "Pack-requirements badge", "Show 'Requires: For Rent + Cottage Living + 3 kits' from parsed pack usage. The differentiator.", "End", "V2", "M", "K2"),
    ("K6", "Marketplace", "'You own all packs' check", "For logged-in browsers, compare listing requirements against their pack ownership.", "End", "V2", "S", "K5"),
    ("K7", "Marketplace", "Changelog summarizer", "Turn reimport diff buckets into update copy ('+3 households, 5 lots rebuilt'). Diff engine already exists.", "Super", "V2", "M", ""),
    ("K8", "Marketplace", "Latest-updates feed", "Feed driven by changelogs; bumps freshly-updated listings.", "Super", "V2", "M", "K7,K3"),
    ("K9", "Marketplace", "Moderation + quality gate", "Report/flag, content policy, require cover + % complete to list.", "Dev", "V2", "M", "K3"),
    ("K10", "Marketplace", "SEO on public pages", "Meta tags, OG images, sitemap for showcases + listings.", "Dev", "V2", "M", ""),

    # --- Epic LO: Launch ops (emerged during the 2026-06-03 production push) ---
    ("LO1", "Launch ops", "Confirm Safe Browsing flag cleared", "Google flagged mysavefile.com/api/auth/google as 'Deceptive site' (new-domain OAuth false positive). Review requested in Search Console 2026-06-03. Re-check Security Issues until it shows 'No issues detected' AND the red interstitial is gone in-browser. Gates ONLY Google sign-in; email/password login is unaffected.", "Dev", "P0", "S", ""),
    ("LO2", "Launch ops", "Production domain + email + secrets", "mysavefile.com on Cloudflare DNS + Railway (HTTPS), Resend email from mail.mysavefile.com, R2 keys rotated to a least-privilege token, privacy@ via Cloudflare Email Routing, legal pages linked everywhere. Full detail in DEPLOYMENT.md / project_production_infra memory.", "Dev", "P0", "L", ""),
    ("LO3", "Launch ops", "Populate landing spotlight (feature a real save)", "On prod DB, feature at least one public showcase so the landing's 'See a real save' spotlight isn't hidden. The save must have a share_token (public sharing on). Run: UPDATE save_files SET featured_at = NOW() WHERE name = '<your save>' AND share_token IS NOT NULL;  (set = NULL to un-feature). See project_landing_and_brand memory.", "Dev", "P0", "S", "B3,LO2"),

    # --- Epic L: Parser completeness / "full tool" (auto-from-save) ---
    # Feasibility spiked & PROVEN 2026-06-03/04 (memories: project_parser_completeness,
    # project_lineage_spike, project_dynasty_spike). Everything auto-parses from the
    # .save, preserving the import-magic. DECISION 2026-06-04: sequence this lane
    # BEFORE Epic F (mobile) — settle the core information architecture of the new
    # pages/links before finalizing mobile, to avoid backtracking mobile work.
    # CAVEAT: big lane — start with L1 (IA) + L2 (lineage); don't let it indefinitely
    # delay launch. NOTE: household bios already auto-extracted (households.ts f18).
    ("L1", "Parser completeness", "Information architecture for new surfaces", "Decide where lineage / dynasties / careers live (new pages vs tabs on Sims/Households), public-vs-creator-only, and how they link into existing pages (Sims roster, household cards, showcase). Cheap design pass that gates the rest. Also verify lot + dynasty description fields are parsed.", "Dev", "P1", "S", ""),
    ("L2", "Parser completeness", "Family tree / lineage", "Parse f15 (spouse) + f30 (blood+marriage family incl. siblings/in-laws, cross-household) into a family-tree view. IA (L1): a DEDICATED Family Trees page under Tools PLUS inline 'view family tree' entry points from sim rows / household cards / showcase (L1a=Both). Keep the tree renderer reusable for public/creator-preview later. Non-marital romance out of scope (relationship-track data).", "Super", "P1", "L", "L1"),
    ("L3", "Parser completeness", "Careers on the Sims roster", "Parse the career_tracker (field 30; skills live in field 21 = deferred) + build a career-uid->name map (buildStock* pattern). IA (L1): add a Career column to the Sims roster (sortable, beside traits/aspirations) + a Careers tab in Diversity for the distribution. No new nav item.", "Super", "P1", "M", "L1"),
    ("L4", "Parser completeness", "Dynasties view", "Parse the in-game Dynasty record from 0x0d (name, member roster, roles Head/Heir/Member/Outcast + succession order, progression + royal/commoner favor, alliances/rivalries). IA (L1): a Dynasties page under Manage (pack-gated, beside Clubs); roster cross-links into the Family Trees view + member sims. Feasible per project_dynasty_spike.", "Super", "P1", "L", "L1"),
    ("L5", "Parser completeness", "Custom-venue scheduler (Tier 1-2)", "Detect a lot's active venue schedule: custom preset/role NAMES (stored as free strings), role counts, timeslot times. Skip deep criteria/activities (Tier 3). App already supports hand-entered venue_schedule as fallback. See project_parser_completeness.", "Super", "P1", "M", "L1"),
    ("L6", "Parser completeness", "Custom-venue scheduler (Tier 3, deep)", "Full role criteria/activities/outfit + Maxis-preset expansion from game packages + tuning->name maps. Hardest + most niche; only if creators ask.", "Super", "P2", "L", "L5"),

    # --- Epic PH: Photos & media (surfaced during dogfooding 2026-06-08) ---
    ("PH1", "Photos & media", "Inspo photo load issue", "Inspo gallery photos fail to load / load unreliably [reported during dogfooding 2026-06-08 — needs repro + root cause]. Investigate load order, R2 fetch, and lazy-loading in the inspo gallery.", "End", "P1", "M", ""),
    ("PH2", "Photos & media", "Multiple showcase photos per lot — display UX", "A lot can hold unlimited showcase photos (no cap server-side); the showcase lightbox/gallery isn't paginated or virtualized and the lot card surfaces only a single cover. Decide how to present many photos on one lot: gallery grid, ordering, primary-cover selection, lazy-loading. (Uploads now auto-compress to ~0.7MB JPEG, so size is no longer the concern — presentation is.)", "End", "P1", "M", ""),
]

# Live status overrides keyed by task ID — update as work completes.
STATUS_OVERRIDES = {
    "A1": "Done", "A2": "Done", "A3": "Done", "A4": "Done",
    "A5": "Done", "A6": "Done", "A7": "Done",
    "C1": "Done", "C2": "Done", "C3": "Done", "C5": "Done",
    "D1": "Done", "D2": "Done", "D3": "Done", "D5": "Done",
    "D4": "Deferred - no fetch caller yet (guard ready)",
    "E1": "Done", "E2": "Done", "E3": "Done", "E4": "Done",
    "A8": "Done",
    "B1": "Done", "B2": "Done", "B5": "Done", "B6": "Done",
    "B3": "Done - single real-save spotlight (scales to strip at n>=2)",
    "B4": "Done - featured_at column, curated via SQL for now",
    "G1": "Done - boot-status logs; verify DSNs in Railway",
    "G2": "Done - vite plugin gated on SENTRY_AUTH_TOKEN",
    "G3": "Done - documented; create alert rule in Sentry dashboard",
    "G4": "Done - login 15/15min, register 10/hr, public 300/15min per IP",
    "G5": "Deferred - revisit post-launch (no users to measure yet)",
    "H1": "Done - 21 protobuf + parseClub regression tests added",
    "H2": "Documented - graceful 'Stock Club' fallback + Help note + guard test",
    "I1": "Done - description column + plumbing; import writes description not notes",
    "I2": "Done - per-world showcase blurb (world_blurbs), editable in WorldShowcase",
    "I3": "Done - labeled Description (public) + Notes (private) in all 4 editors",
    "I4": "Done - descriptions + world blurbs on public showcase; plugged a notes leak",
    "I5": "Done - traits/aspirations + sortable columns on the Sims roster",
    "I6": "CUT - .save stays the single source of truth; a 2nd import path muddies it",
    "LO2": "Done",
    "LO1": "Pending - re-check Search Console",
    "LO3": "Pending - feature a real save on prod after deploy",
    "L1": "Done 2026-06-05 - IA decided: careers=Sims column+Diversity tab; family trees=dedicated /family page (Tools) + inline entry from sims/households/showcase (Both); dynasties=/dynasties under Manage (pack-gated); slot into existing nav groups, no reorg.",
}

PRIORITY_FILL = {
    "P0": "F4CCCC",  # red-ish
    "P1": "FCE5CD",  # amber
    "P2": "EAEAEA",  # grey
    "GW": "D9EAD3",  # green (groundwork)
    "V2": "E6D9F2",  # plum (v2)
}
HEADER_FILL = "7C5CBF"  # brand plum
COLS = ["ID", "Epic", "Task", "Detail", "Persona", "Priority", "Size", "Deps", "Status"]
WIDTHS = [6, 22, 30, 60, 9, 9, 6, 10, 12]


def style_sheet(ws, rows):
    thin = Side(style="thin", color="DDDDDD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    # header
    for c, name in enumerate(COLS, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(vertical="center", horizontal="left")
        cell.border = border
        ws.column_dimensions[get_column_letter(c)].width = WIDTHS[c - 1]
    ws.row_dimensions[1].height = 22
    # body
    for r, row in enumerate(rows, start=2):
        # row is (ID,Epic,Task,Detail,Persona,Priority,Size,Deps); Status from overrides
        values = list(row[:8]) + [STATUS_OVERRIDES.get(row[0], "")]
        prio = row[5]
        for c, val in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.alignment = Alignment(vertical="top", wrap_text=(c in (2, 3, 4)))
            cell.border = border
            cell.font = Font(size=10)
            if c == 6:  # priority column gets the fill + bold
                cell.fill = PatternFill("solid", fgColor=PRIORITY_FILL.get(prio, "FFFFFF"))
                cell.font = Font(size=10, bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{len(rows)+1}"


def build_xlsx(path):
    wb = Workbook()
    launch = [t for t in TASKS if t[5] in ("P0", "P1", "P2")]
    market = [t for t in TASKS if t[5] in ("GW", "V2")]

    ws1 = wb.active
    ws1.title = "Launch Plan"
    style_sheet(ws1, launch)

    ws2 = wb.create_sheet("Marketplace v2")
    style_sheet(ws2, market)

    ws3 = wb.create_sheet("Legend")
    legend = [
        ("Go-live bar", "A stranger can sign up, get their save (or a starter) in front of them, NOT lose data, and recover their account."),
        ("", ""),
        ("Priority", ""),
        ("P0", "Go-live gate. Launch blocks on these."),
        ("P1", "Strongly want at launch; not an absolute blocker."),
        ("P2", "Fast-follow after launch."),
        ("GW", "v2 groundwork - cheap to do at launch so v2 is a query, not a rewrite."),
        ("V2", "Marketplace v2 - flagship post-launch, its own focused build."),
        ("", ""),
        ("Persona", ""),
        ("Dev", "Developer / Product Manager - reliability, safety, legal, security."),
        ("Super", "Sims super user - migration in, parity, self-marketing."),
        ("End", "Sims end user - ease, fun, mobile, browse without building."),
        ("", ""),
        ("Size", "S = small, M = medium, L = large."),
        ("", ""),
        ("Note", "Single source of truth is scripts/diagnostics/genLaunchPlan.py. Edit + re-run to regenerate xlsx + md."),
    ]
    for r, (k, v) in enumerate(legend, start=1):
        a = ws3.cell(row=r, column=1, value=k)
        a.font = Font(bold=True, color="7C5CBF" if k in ("Go-live bar", "Priority", "Persona", "Size", "Note") else "000000")
        b = ws3.cell(row=r, column=2, value=v)
        b.alignment = Alignment(wrap_text=True, vertical="top")
    ws3.column_dimensions["A"].width = 16
    ws3.column_dimensions["B"].width = 90

    wb.save(path)
    return len(launch), len(market)


def build_md(path):
    by_prio = {"P0": [], "P1": [], "P2": [], "GW": [], "V2": []}
    for t in TASKS:
        by_prio[t[5]].append(t)

    def table(rows):
        out = ["| ID | Epic | Task | Detail | Persona | Size | Deps | Status |",
               "|---|---|---|---|---|---|---|---|"]
        for t in rows:
            out.append(f"| {t[0]} | {t[1]} | {t[2]} | {t[3]} | {t[4]} | {t[6]} | {t[7] or '-'} | {STATUS_OVERRIDES.get(t[0], '')} |")
        return "\n".join(out)

    lines = []
    lines.append("# Launch Plan — Sims Save File Planner")
    lines.append("")
    lines.append("> Single source of truth: `scripts/diagnostics/genLaunchPlan.py`. Edit a task there and re-run to regenerate this file and `LAUNCH_PLAN.xlsx`.")
    lines.append("")
    lines.append("**Go-live bar:** a stranger can sign up, get their save (or a starter) in front of them, not lose data, and recover their account.")
    lines.append("")
    lines.append("**Priority:** P0 = go-live gate · P1 = want at launch · P2 = fast-follow · GW = v2 groundwork (do at launch) · V2 = marketplace v2.")
    lines.append("**Persona:** Dev = Developer/PM · Super = Sims super user · End = Sims end user.")
    lines.append("")
    lines.append("## P0 — go-live gates")
    lines.append("")
    lines.append(table(by_prio["P0"]))
    lines.append("")
    lines.append("## P1 — want at launch")
    lines.append("")
    lines.append(table(by_prio["P1"]))
    lines.append("")
    lines.append("## P2 — fast-follow")
    lines.append("")
    lines.append(table(by_prio["P2"]))
    lines.append("")
    lines.append("## Marketplace — groundwork at launch (GW)")
    lines.append("")
    lines.append(table(by_prio["GW"]))
    lines.append("")
    lines.append("## Marketplace v2 — post-launch flagship (V2)")
    lines.append("")
    lines.append(table(by_prio["V2"]))
    lines.append("")
    with open(path, "w") as f:
        f.write("\n".join(lines))
    return {k: len(v) for k, v in by_prio.items()}


if __name__ == "__main__":
    xlsx_path = os.path.join(ROOT, "LAUNCH_PLAN.xlsx")
    md_path = os.path.join(ROOT, "LAUNCH_PLAN.md")
    nl, nm = build_xlsx(xlsx_path)
    counts = build_md(md_path)
    print(f"xlsx: {xlsx_path}  (Launch Plan={nl}, Marketplace v2={nm})")
    print(f"md:   {md_path}")
    print("counts:", counts, "total:", len(TASKS))
